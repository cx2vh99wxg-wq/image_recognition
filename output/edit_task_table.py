from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
src=r'C:/Users/LENOVO/Desktop/技术文档模板(1).docx'
out=r'D:/vmshare/image_recognition/output/技术文档模板_任务表修订.docx'
d=Document(src);t=d.tables[0]
rows=[
('基础任务 1\n图像采集与显示','通过 HDMI、摄像头等输入接口采集图像，经 LCD 或 HDMI 输出，完成采集—处理—显示闭环，实现道路环境感知。','六路 OV5640 采集周边画面，完成图像处理及 HDMI 多画面显示。'),
('基础任务 2\n软硬协同通信','设计并实现 MCU 与 FPGA 之间的高速通信机制（如 PCIe），完成数据交互与指令控制。','板内经 PCIe 传输图像、FSPI 传递控制信息；双板经 UDP 交互，进程使用共享内存。'),
('基础任务 3\n图像处理与识别','基于 FPGA 实现系统对道路周围环境的监测和特定场景的感知。','FPGA 完成视频采集及一路 CNN 夜视增强；结合 ARM 端识别完成道路场景感知。'),
('高阶任务 1\n多路采集与传输','利用 FPGA 与 ARM 高速通信接口，实现至少四路视频流并行采集与低延迟传输，并优化数据流带宽。','双板各采集三路，共六路；采用 DDR3 缓存、PCIe DMA 和 UDP 传输。端到端延迟待测。'),
('高阶任务 2\n复杂应用场景','针对夜间道路、复杂交通路口等特定场景，构建完整的感知和避障功能。','实现单路夜视增强、信号灯及行人识别和制动控制；复杂场景避障闭环待验证。'),
('高阶任务 3\nAI 部署与识别','在处理器端部署轻量级 AI 模型，实现多路视频实时、高精度、低延迟识别及自主道路规划和辅助驾驶。','M 端运行 YOLOPv2、YOLOv5n；S 端运行 YOLOv5s。支持车道感知和辅助控制；自主规划及性能待验证。'),
('补充 1\n信号灯场景','识别信号灯，根据灯色约束车辆行驶。','实现红灯前进限制与绿灯解除限制逻辑；场景联动及报警待验证。'),
('补充 2\n车道线与变道','根据实、虚线限制或允许变道并提示。','确认实线时拦截人工转向请求并提示；完整自动变道待验证。'),
('补充 3\n行人制动','识别行人等危险目标并实施制动。','行人检测结果经共享内存、FSPI 触发 FPGA 制动锁存；实车效果待测。'),
('补充 4\n人机交互','实现车辆操作及状态信息呈现。','HMI 串口屏输入控制指令；HDMI 展示画面及识别信息。'),
('补充 5\n实时性与性能','测试视频帧率、延迟及识别效果。','10 月 25 日记录：M 发布约 30 fps、S 接收约 30 fps、显示渲染约 12 fps；最新版本待复测。'),
('补充 6\n导航与规划','结合导航与道路场景实现模拟智驾。','具备弯道信息驱动转向链路；导航驱动的自主规划尚待验证。')]
for ri,values in enumerate([('项目','得分要求','完成情况')]+rows):
 for ci,(cell,value) in enumerate(zip(t.rows[ri].cells,values)):
  cell.text=value;cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cell.width=Cm([2.6,5.8,6.2][ci]);tcPr=cell._tc.get_or_add_tcPr();shd=tcPr.find(qn('w:shd'))
  if shd is None: shd=OxmlElement('w:shd');tcPr.append(shd)
  shd.set(qn('w:fill'),'D9D9D9' if ri==0 else ('F4F4F4' if ri%2==0 else 'FFFFFF'))
  for p in cell.paragraphs:
   p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(10.5)
   snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(snap)
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (ci==0 or ri==0) else WD_ALIGN_PARAGRAPH.LEFT
   for run in p.runs:
    run.font.name='宋体';run._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');run.font.size=Pt(8.5);run.bold=(ri==0);run.font.color.rgb=RGBColor(0,0,0)
 trPr=t.rows[ri]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
 for height in trPr.findall(qn('w:trHeight')): trPr.remove(height)
t.autofit=False
for col,w in zip(t.columns,[2.6,5.8,6.2]):col.width=Cm(w)
# Table starts on a fresh page without changing the preceding prose.
p=OxmlElement('w:p');run=OxmlElement('w:r');br=OxmlElement('w:br');br.set(qn('w:type'),'page');run.append(br);p.append(run);t._element.addprevious(p)
d.save(out)
print(out)
